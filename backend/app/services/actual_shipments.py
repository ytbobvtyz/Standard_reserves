from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import APIError
from app.models.actual_shipment import ActualShipment
from app.models.client import Client
from app.models.object import Object
from app.models.product import Product
from app.schemas.actual_shipments import (
    ActualShipmentError,
    ActualShipmentListItem,
    ActualShipmentMissingClient,
    ActualShipmentMissingProduct,
    ActualShipmentPreviewRow,
    ActualShipmentUploadResult,
)

PREVIEW_LIMIT = 50
WEIGHT_UNIT = "кг"
TEMPLATE_HEADERS = (
    "Пункт отгрузки",
    "Клиент",
    "Имя клиента",
    "№ продукта",
    "Назв. позиции",
    "Вес нетто",
    "Дата фактического ОМ",
    "Транспортный заказ",
    "Описание Маршрута",
)
REQUIRED_COLUMNS = (
    "loading_point",
    "client_code",
    "product_code",
    "quantity",
    "shipment_date",
)
REQUIRED_LABELS = {
    "loading_point": "Пункт отгрузки",
    "client_code": "Клиент",
    "product_code": "№ продукта",
    "quantity": "Вес нетто",
    "shipment_date": "Дата фактического ОМ",
}
HEADER_ALIASES = {
    "пункт отгрузки": "loading_point",
    "loading_point": "loading_point",
    "клиент": "client_code",
    "client_code": "client_code",
    "кссс": "client_code",
    "id клиента": "client_code",
    "код клиента": "client_code",
    "имя клиента": "client_name",
    "наименование клиента": "client_name",
    "n продукта": "product_code",
    "номер продукта": "product_code",
    "product_code": "product_code",
    "артикул": "product_code",
    "назв. позиции": "product_name",
    "название позиции": "product_name",
    "наименование продукции": "product_name",
    "вес нетто": "quantity",
    "quantity": "quantity",
    "количество": "quantity",
    "дата фактического ом": "shipment_date",
    "shipment_date": "shipment_date",
    "дата отгрузки": "shipment_date",
    "транспортный заказ": "invoice_number",
    "invoice_number": "invoice_number",
    "накладная": "invoice_number",
    "накладной": "invoice_number",
    "n накладной": "invoice_number",
    "номер накладной": "invoice_number",
    "описание маршрута": "route",
    "route": "route",
    "маршрут": "route",
}
HEADER_HINTS = {
    "A1": "Код пункта отгрузки, до 10 символов. Обязательное.",
    "B1": "КССС клиента. Обязательное. Неизвестные клиенты пропускаются.",
    "C1": "Наименование клиента из ERP. Справочное, в базу не пишется.",
    "D1": "Артикул. Обязательное. Неизвестные продукты пропускаются.",
    "E1": "Наименование продукции из ERP. Справочное, в базу не пишется.",
    "F1": "Вес нетто в килограммах, больше 0. Обязательное.",
    "G1": "Дата ОМ, ДД.ММ.ГГГГ. Пустая ячейка — строка пропускается.",
    "H1": "Транспортный заказ, до 25 символов.",
    "I1": "Описание маршрута, до 50 символов.",
}


@dataclass
class _RawRow:
    row: int
    loading_point: str
    client_code: int
    file_client_name: str | None
    product_code: int
    file_product_name: str | None
    quantity: Decimal
    shipment_date: date
    invoice_number: str | None
    route: str | None


@dataclass
class _ValidRow:
    raw: _RawRow
    client: Client
    product: Product


def _check_period(date_from: date | None, date_to: date | None) -> None:
    if date_from is not None and date_to is not None and date_from > date_to:
        raise APIError(
            400,
            "VALIDATION_ERROR",
            "date_from не может быть позже date_to",
        )


def build_template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Реализация"
    sheet.append(list(TEMPLATE_HEADERS))
    for cell, hint in HEADER_HINTS.items():
        sheet[cell].comment = Comment(hint, "Standart Reserve")
    widths = {
        "A": 18,
        "B": 14,
        "C": 28,
        "D": 16,
        "E": 36,
        "F": 14,
        "G": 24,
        "H": 22,
        "I": 42,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text.lower() in {"none", "nan"}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def _header_key(value: Any) -> str:
    text = _cell_text(value).casefold().replace("ё", "е")
    text = text.replace("№", "n").strip().rstrip(".")
    return " ".join(text.split())


def _header_map(header: tuple[Any, ...]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, cell in enumerate(header):
        key = HEADER_ALIASES.get(_header_key(cell))
        if key and key not in columns:
            columns[key] = index
    return columns


def _cell(row: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns.get(name)
    if index is None or index >= len(row):
        return None
    return row[index]


def _row_empty(row: tuple[Any, ...], columns: dict[str, int]) -> bool:
    names = (
        "loading_point",
        "client_code",
        "client_name",
        "product_code",
        "product_name",
        "quantity",
        "shipment_date",
        "invoice_number",
        "route",
    )
    return all(_is_blank(_cell(row, columns, name)) for name in names)


def _parse_int(value: Any, field: str, *, maximum: int | None = None) -> int:
    if _is_blank(value):
        raise ValueError(f"Укажите {field}")
    if isinstance(value, bool):
        raise ValueError(f"{field} должен быть числом")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float):
        parsed = int(value)
        if float(parsed) != value:
            raise ValueError(f"{field} должен быть целым числом")
    elif isinstance(value, Decimal):
        parsed = int(value)
        if Decimal(parsed) != value:
            raise ValueError(f"{field} должен быть целым числом")
    else:
        text = str(value).strip().replace(" ", "").replace(",", ".")
        if text.endswith(".0"):
            text = text[:-2]
        try:
            as_float = float(text)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{field} должен быть числом") from exc
        parsed = int(as_float)
        if as_float != float(parsed):
            raise ValueError(f"{field} должен быть целым числом")
    if parsed <= 0:
        raise ValueError(f"{field} должен быть больше 0")
    if maximum is not None and parsed > maximum:
        raise ValueError(f"{field} слишком большой")
    return parsed


def _parse_loading_point(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, int) and not isinstance(value, bool):
        text = str(value)
    else:
        text = _cell_text(value)
    if not text:
        raise ValueError("Укажите пункт отгрузки")
    if len(text) > 10:
        raise ValueError("Пункт отгрузки не длиннее 10 символов")
    return text


def _parse_reference_name(value: Any) -> str | None:
    if _is_blank(value):
        return None
    text = str(value).strip()
    if len(text) > 200:
        return text[:200]
    return text


def _parse_quantity(value: Any) -> Decimal:
    if _is_blank(value):
        raise ValueError("Укажите вес нетто")
    if isinstance(value, bool):
        raise ValueError("Вес нетто должен быть числом")
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
    else:
        text = str(value).strip().replace(" ", "").replace(",", ".")
        try:
            amount = Decimal(text)
        except (InvalidOperation, TypeError, ValueError) as exc:
            raise ValueError("Вес нетто должен быть числом") from exc
    if amount <= 0:
        raise ValueError("Вес нетто должен быть больше 0")
    if amount >= Decimal("10000000000"):
        raise ValueError("Вес нетто слишком большой")
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if _is_blank(value):
        raise ValueError("Укажите дату фактического ОМ")
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError("Дата фактического ОМ должна быть в формате ДД.ММ.ГГГГ")


def _parse_optional(value: Any, field: str, limit: int) -> str | None:
    if _is_blank(value):
        return None
    text = str(value).strip()
    if len(text) > limit:
        raise ValueError(f"{field} не длиннее {limit} символов")
    return text


def _load_workbook(content: bytes, filename: str):
    lowered = filename.lower()
    if not lowered.endswith((".xlsx", ".xls")):
        raise APIError(400, "INVALID_FILE", "Загрузите файл .xlsx или .xls")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise APIError(
            400,
            "INVALID_FILE",
            "Не удалось прочитать Excel. Используйте формат .xlsx",
        ) from exc
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise APIError(400, "INVALID_FILE", "Файл пуст")
    return rows


def _parse_rows(
    rows: list[tuple[Any, ...]],
) -> tuple[list[_RawRow], list[ActualShipmentError], int, int]:
    columns = _header_map(rows[0])
    missing = [
        REQUIRED_LABELS[name] for name in REQUIRED_COLUMNS if name not in columns
    ]
    if missing:
        raise APIError(
            400,
            "INVALID_FILE",
            "В файле нет колонок: " + ", ".join(missing),
        )
    parsed: list[_RawRow] = []
    errors: list[ActualShipmentError] = []
    total = 0
    skipped = 0
    for excel_row, row in enumerate(rows[1:], start=2):
        if row is None or _row_empty(row, columns):
            continue
        total += 1
        if _is_blank(_cell(row, columns, "shipment_date")):
            skipped += 1
            continue
        try:
            parsed.append(
                _RawRow(
                    row=excel_row,
                    loading_point=_parse_loading_point(
                        _cell(row, columns, "loading_point")
                    ),
                    client_code=_parse_int(
                        _cell(row, columns, "client_code"),
                        "Клиент",
                        maximum=9999999999,
                    ),
                    file_client_name=_parse_reference_name(
                        _cell(row, columns, "client_name")
                    ),
                    product_code=_parse_int(
                        _cell(row, columns, "product_code"),
                        "№ продукта",
                    ),
                    file_product_name=_parse_reference_name(
                        _cell(row, columns, "product_name")
                    ),
                    quantity=_parse_quantity(_cell(row, columns, "quantity")),
                    shipment_date=_parse_date(_cell(row, columns, "shipment_date")),
                    invoice_number=_parse_optional(
                        _cell(row, columns, "invoice_number"),
                        "Транспортный заказ",
                        25,
                    ),
                    route=_parse_optional(
                        _cell(row, columns, "route"),
                        "Описание маршрута",
                        50,
                    ),
                )
            )
        except ValueError as exc:
            errors.append(ActualShipmentError(row=excel_row, message=str(exc)))
    return parsed, errors, total, skipped


async def _clients_by_code(db: AsyncSession, codes: set[int]) -> dict[int, Client]:
    if not codes:
        return {}
    result = await db.execute(
        select(Client).where(
            Client.client_code.in_(codes),
            Client.deleted_at.is_(None),
        )
    )
    return {item.client_code: item for item in result.scalars().all()}


async def _products_by_code(db: AsyncSession, codes: set[int]) -> dict[int, Product]:
    if not codes:
        return {}
    result = await db.execute(
        select(Product).where(
            Product.code.in_(codes),
            Product.deleted_at.is_(None),
        )
    )
    return {item.code: item for item in result.scalars().all()}


def _preview_row(item: _ValidRow) -> ActualShipmentPreviewRow:
    raw = item.raw
    return ActualShipmentPreviewRow(
        row=raw.row,
        client_id=item.client.id,
        client_code=item.client.client_code,
        client_name=item.client.name,
        product_code=item.product.code,
        product_name=item.product.name,
        quantity=raw.quantity,
        unit=WEIGHT_UNIT,
        shipment_date=raw.shipment_date,
        loading_point=raw.loading_point,
        invoice_number=raw.invoice_number,
        route=raw.route,
    )


def _result_message(
    *,
    saved: bool,
    valid_rows: int,
    errors: int,
    inserted: int,
    skipped_rows: int,
    missing_clients: int,
    missing_products: int,
) -> str:
    notes: list[str] = []
    if skipped_rows:
        notes.append(f"без даты отгрузки {skipped_rows}")
    if missing_clients:
        notes.append(f"клиентов нет в справочнике {missing_clients}")
    if missing_products:
        notes.append(f"продуктов нет в справочнике {missing_products}")
    if saved and not notes and errors == 0:
        return f"Сохранено {inserted}, ошибок 0"
    tail = ", ".join(notes)
    if saved:
        saved_text = f"Сохранено {inserted}, ошибок {errors}"
        return f"{saved_text}. {tail}." if tail else saved_text
    shown = min(valid_rows, PREVIEW_LIMIT)
    parts = [f"К загрузке {valid_rows} строк", f"ошибок {errors}"]
    if tail:
        parts.insert(1, tail)
    text = ", ".join(parts)
    if valid_rows > PREVIEW_LIMIT:
        return f"{text}. В предпросмотре первые {shown}. Данные не сохранены."
    return f"{text}. Данные не сохранены."


async def upload_shipments(
    db: AsyncSession,
    content: bytes,
    filename: str,
    *,
    user_id: UUID,
    confirm: bool,
) -> ActualShipmentUploadResult:
    rows = _load_workbook(content, filename)
    parsed, errors, total, skipped_rows = _parse_rows(rows)
    clients = await _clients_by_code(db, {item.client_code for item in parsed})
    products = await _products_by_code(db, {item.product_code for item in parsed})
    valid: list[_ValidRow] = []
    missing_client_rows: dict[int, ActualShipmentMissingClient] = {}
    missing_product_rows: dict[int, ActualShipmentMissingProduct] = {}
    for item in parsed:
        client = clients.get(item.client_code)
        product = products.get(item.product_code)
        if client is None:
            known = missing_client_rows.get(item.client_code)
            if known is None:
                missing_client_rows[item.client_code] = ActualShipmentMissingClient(
                    client_code=item.client_code,
                    client_name=item.file_client_name,
                    rows=1,
                )
            else:
                known.rows += 1
        if product is None:
            known_product = missing_product_rows.get(item.product_code)
            if known_product is None:
                missing_product_rows[item.product_code] = ActualShipmentMissingProduct(
                    product_code=item.product_code,
                    product_name=item.file_product_name,
                    rows=1,
                )
            else:
                known_product.rows += 1
        if client is None or product is None:
            continue
        valid.append(_ValidRow(raw=item, client=client, product=product))
    errors.sort(key=lambda item: item.row)
    missing_clients = sorted(
        missing_client_rows.values(),
        key=lambda item: item.client_code,
    )
    missing_products = sorted(
        missing_product_rows.values(),
        key=lambda item: item.product_code,
    )

    inserted = 0
    if confirm and valid:
        for item in valid:
            raw = item.raw
            db.add(
                ActualShipment(
                    client_id=item.client.id,
                    product_code=item.product.code,
                    quantity=raw.quantity,
                    unit=WEIGHT_UNIT,
                    shipment_date=raw.shipment_date,
                    loading_point=raw.loading_point,
                    invoice_number=raw.invoice_number,
                    route=raw.route,
                    uploaded_by=user_id,
                )
            )
        await db.commit()
        inserted = len(valid)

    saved = confirm and inserted > 0
    return ActualShipmentUploadResult(
        saved=saved,
        total_rows=total,
        valid_rows=len(valid),
        inserted=inserted,
        skipped_rows=skipped_rows,
        errors=len(errors),
        message=_result_message(
            saved=saved,
            valid_rows=len(valid),
            errors=len(errors),
            inserted=inserted,
            skipped_rows=skipped_rows,
            missing_clients=len(missing_clients),
            missing_products=len(missing_products),
        ),
        preview=[_preview_row(item) for item in valid[:PREVIEW_LIMIT]],
        error_details=errors,
        missing_clients=missing_clients,
        missing_products=missing_products,
    )


def _to_list_item(row: ActualShipment) -> ActualShipmentListItem:
    return ActualShipmentListItem(
        id=row.id,
        client_id=row.client_id,
        client_code=row.client.client_code,
        client_name=row.client.name,
        product_code=row.product_code,
        product_name=row.product.name,
        quantity=row.quantity,
        unit=row.unit,  # type: ignore[arg-type]
        shipment_date=row.shipment_date,
        loading_point=row.loading_point,
        invoice_number=row.invoice_number,
        route=row.route,
    )


async def list_shipments(
    db: AsyncSession,
    *,
    client_id: UUID | None = None,
    product_code: int | None = None,
    warehouse_code: int | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = 1,
    limit: int = 50,
) -> tuple[list[ActualShipmentListItem], int]:
    _check_period(date_from, date_to)
    conditions = [ActualShipment.deleted_at.is_(None)]
    if warehouse_code is not None:
        loading_point = await db.scalar(
            select(Object.loading_point).where(
                Object.code == warehouse_code,
                Object.deleted_at.is_(None),
            )
        )
        if loading_point is None:
            return [], 0
        conditions.append(ActualShipment.loading_point == loading_point)
    if client_id is not None:
        conditions.append(ActualShipment.client_id == client_id)
    if product_code is not None:
        conditions.append(ActualShipment.product_code == product_code)
    if date_from is not None:
        conditions.append(ActualShipment.shipment_date >= date_from)
    if date_to is not None:
        conditions.append(ActualShipment.shipment_date <= date_to)

    total = int(
        await db.scalar(
            select(func.count()).select_from(ActualShipment).where(*conditions)
        )
        or 0
    )
    result = await db.execute(
        select(ActualShipment)
        .options(
            selectinload(ActualShipment.client),
            selectinload(ActualShipment.product),
        )
        .where(*conditions)
        .order_by(
            ActualShipment.shipment_date.desc(),
            ActualShipment.created_at.desc(),
            ActualShipment.id,
        )
        .offset((page - 1) * limit)
        .limit(limit)
    )
    return [_to_list_item(row) for row in result.scalars().all()], total
