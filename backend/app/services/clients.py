import re
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import APIError
from app.core.pagination import paginate
from app.models.client import (
    SYSTEM_CLIENT_CODE,
    SYSTEM_CLIENT_ID,
    SYSTEM_CLIENT_INN,
    Client,
)
from app.models.department import Department
from app.models.user import User
from app.schemas.client import (
    ClientCreate,
    ClientListItem,
    ClientUpdate,
    ClientUploadDepartmentPlan,
    ClientUploadError,
    ClientUploadPreview,
    ClientUploadResult,
)
from app.schemas.common import PaginationMeta

INN_RE = re.compile(r"^\d{10}$|^\d{12}$")
KPP_RE = re.compile(r"^\d{9}$")
TEMPLATE_HEADERS = (
    "ID клиента (КССС)",
    "Наименование",
    "ИНН",
    "КПП",
    "ID подразделения",
)
EXPORT_HEADERS = (
    *TEMPLATE_HEADERS,
    "Управление продаж",
    "Статус",
)
HEADER_ALIASES = {
    "id клиента (кссс)": "client_code",
    "id клиента": "client_code",
    "кссс": "client_code",
    "client_code": "client_code",
    "наименование": "name",
    "название": "name",
    "name": "name",
    "инн": "inn",
    "inn": "inn",
    "кпп": "kpp",
    "kpp": "kpp",
    "id подразделения": "department",
    "подразделение": "department",
    "department_id": "department",
    "управление продаж": "department_name",
    "название подразделения": "department_name",
    "department_name": "department_name",
}
NO_DEPARTMENT = "Без управления"
DEPARTMENT_HINT = (
    'Укажите название подразделения (например, "Коммерческий отдел") или его UUID'
)


def names_match(left: str | None, right: str | None) -> bool:
    return (left or "").strip().casefold() == (right or "").strip().casefold()


def to_list_item(client: Client) -> ClientListItem:
    department = client.department
    return ClientListItem(
        id=client.id,
        client_code=client.client_code,
        inn=client.inn,
        name=client.name,
        kpp=client.kpp,
        department_id=client.department_id,
        department_name=department.name if department is not None else None,
        is_active=client.is_active,
        is_system=client.is_system,
    )


def _search_clause(query: str | None):
    if query is None or not query.strip():
        return None
    term = query.strip()
    if term.isdigit():
        clauses = [Client.name.ilike(f"%{term}%")]
        if 1 <= len(term) <= 10:
            clauses.append(Client.client_code == int(term))
        if len(term) in {10, 12}:
            clauses.append(Client.inn == term)
        elif len(term) < 10:
            clauses.append(Client.inn.like(f"{term}%"))
        return or_(*clauses)
    return Client.name.ilike(f"%{term}%")


async def list_clients(
    db: AsyncSession,
    *,
    q: str | None,
    department_id: UUID | None,
    is_active: bool | None,
    include_inactive: bool,
    include_deleted: bool,
    page: int,
    limit: int,
) -> tuple[list[ClientListItem], PaginationMeta]:
    conditions = []
    if not include_deleted:
        conditions.append(Client.deleted_at.is_(None))
    if not include_inactive:
        active = True if is_active is None else is_active
        conditions.append(Client.is_active.is_(active))
    elif is_active is not None:
        conditions.append(Client.is_active.is_(is_active))
    if department_id is not None:
        conditions.append(Client.department_id == department_id)
    search = _search_clause(q)
    if search is not None:
        conditions.append(search)
    count_stmt = select(func.count()).select_from(Client)
    list_stmt = select(Client).options(selectinload(Client.department))
    if conditions:
        count_stmt = count_stmt.where(*conditions)
        list_stmt = list_stmt.where(*conditions)
    total = await db.scalar(count_stmt)
    result = await db.execute(
        paginate(list_stmt.order_by(Client.name, Client.client_code), page, limit)
    )
    items = [to_list_item(client) for client in result.scalars().all()]
    return items, PaginationMeta(page=page, limit=limit, total=int(total or 0))


async def get_client(db: AsyncSession, client_id: UUID) -> Client:
    result = await db.execute(
        select(Client)
        .options(selectinload(Client.department))
        .where(Client.id == client_id, Client.deleted_at.is_(None))
    )
    client = result.scalar_one_or_none()
    if client is None:
        raise APIError(404, "NOT_FOUND", "Клиент не найден")
    return client


async def require_client(db: AsyncSession, client_id: UUID) -> Client:
    return await get_client(db, client_id)


def _validate_code(client_code: int) -> None:
    if client_code < 1 or client_code > 9999999998:
        raise APIError(
            400,
            "VALIDATION_ERROR",
            "ID клиента — целое число от 1 до 10 знаков",
        )
    if client_code == SYSTEM_CLIENT_CODE:
        raise APIError(
            400,
            "VALIDATION_ERROR",
            "Код 9999999999 зарезервирован",
        )


def _validate_inn(inn: str) -> str:
    value = inn.strip()
    if not INN_RE.fullmatch(value) or value == SYSTEM_CLIENT_INN:
        raise APIError(
            400,
            "VALIDATION_ERROR",
            "ИНН должен содержать 10 или 12 цифр",
        )
    return value


def _validate_kpp(kpp: str | None) -> str | None:
    if kpp is None or not kpp.strip():
        return None
    value = kpp.strip()
    if not KPP_RE.fullmatch(value):
        raise APIError(400, "VALIDATION_ERROR", "КПП должен содержать 9 цифр")
    return value


def _validate_name(name: str) -> str:
    value = name.strip()
    if not value or len(value) > 500:
        raise APIError(
            400,
            "VALIDATION_ERROR",
            "Укажите наименование до 500 символов",
        )
    return value


async def _require_department(db: AsyncSession, department_id: UUID) -> Department:
    department = await db.get(Department, department_id)
    if department is None or department.deleted_at is not None:
        raise APIError(400, "VALIDATION_ERROR", "Подразделение не найдено")
    return department


async def _duplicate_code(db: AsyncSession, client_code: int) -> None:
    existing = await db.scalar(
        select(Client).where(
            Client.client_code == client_code,
            Client.deleted_at.is_(None),
        )
    )
    if existing is not None:
        raise APIError(
            409,
            "CONFLICT",
            f"Клиент с кодом {client_code} уже есть ({existing.id})",
        )


async def create_client(
    db: AsyncSession, user: User, body: ClientCreate
) -> ClientListItem:
    _validate_code(body.client_code)
    inn = _validate_inn(body.inn)
    name = _validate_name(body.name)
    kpp = _validate_kpp(body.kpp)
    if body.department_id_set:
        if user.role != "logistics":
            raise APIError(
                403,
                "FORBIDDEN",
                "Подразделение может задать только логист",
            )
        department_id = body.department_id
        if department_id is not None:
            await _require_department(db, department_id)
    else:
        if user.department_id is None:
            raise APIError(
                400,
                "VALIDATION_ERROR",
                "У пользователя не задано подразделение",
            )
        department_id = user.department_id
    await _duplicate_code(db, body.client_code)
    client = Client(
        client_code=body.client_code,
        inn=inn,
        name=name,
        kpp=kpp,
        department_id=department_id,
        is_active=True,
        is_system=False,
    )
    db.add(client)
    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise APIError(
            409,
            "CONFLICT",
            f"Клиент с кодом {body.client_code} уже есть",
        ) from exc
    return to_list_item(await get_client(db, client.id))


async def update_client(
    db: AsyncSession, client_id: UUID, body: ClientUpdate
) -> ClientListItem:
    client = await get_client(db, client_id)
    if client.is_system or client.id == SYSTEM_CLIENT_ID:
        raise APIError(400, "VALIDATION_ERROR", "Служебного клиента нельзя изменить")
    if body.department_id is not None:
        await _require_department(db, body.department_id)
    client.name = _validate_name(body.name)
    client.kpp = _validate_kpp(body.kpp)
    client.is_active = body.is_active
    client.department_id = body.department_id
    client.updated_at = datetime.now(UTC)
    await db.commit()
    return to_list_item(await get_client(db, client.id))


async def delete_client(db: AsyncSession, client_id: UUID) -> None:
    client = await get_client(db, client_id)
    if client.is_system or client.id == SYSTEM_CLIENT_ID:
        raise APIError(400, "VALIDATION_ERROR", "Служебного клиента нельзя удалить")
    client.deleted_at = datetime.now(UTC)
    client.is_active = False
    client.updated_at = datetime.now(UTC)
    await db.commit()


async def system_client(db: AsyncSession) -> Client:
    client = await db.get(Client, SYSTEM_CLIENT_ID)
    if client is None:
        raise APIError(500, "NOT_FOUND", "Служебный клиент не найден")
    return client


async def client_by_code(db: AsyncSession, client_code: int) -> Client | None:
    return await db.scalar(
        select(Client).where(
            Client.client_code == client_code,
            Client.deleted_at.is_(None),
            Client.is_active.is_(True),
        )
    )


def build_template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Клиенты"
    sheet.append(TEMPLATE_HEADERS)
    sheet["E1"].comment = Comment(DEPARTMENT_HINT, "Standart Reserve")
    widths = {"A": 24, "B": 36, "C": 16, "D": 14, "E": 42}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_map(header: tuple[object, ...]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, cell in enumerate(header):
        key = HEADER_ALIASES.get(_cell_text(cell).casefold())
        if key and key not in columns:
            columns[key] = index
    return columns


class _DepartmentIndex:
    def __init__(self, departments: list[Department]) -> None:
        self.by_id = {item.id: item for item in departments}
        self.by_name: dict[str, list[Department]] = {}
        for item in departments:
            key = item.name.strip().lower()
            self.by_name.setdefault(key, []).append(item)

    def from_id_or_name(self, raw: str) -> Department:
        text = raw.strip()
        try:
            department_id = UUID(text)
        except ValueError:
            department_id = None
        if department_id is not None:
            department = self.by_id.get(department_id)
            if department is None:
                raise ValueError(f"Подразделение с UUID {department_id} не найдено")
            return department
        return self.by_exact_name(text)

    def by_exact_name(self, raw: str) -> Department:
        text = raw.strip()
        matches = self.by_name.get(text.lower(), [])
        if len(matches) == 1:
            return matches[0]
        if not matches:
            raise ValueError(f"Подразделение «{text}» не найдено")
        raise ValueError(f"Найдено несколько подразделений с названием «{text}»")


def _resolve_department(
    index: _DepartmentIndex, department_raw: str, department_name_raw: str
) -> Department | None:
    if department_raw.strip():
        return index.from_id_or_name(department_raw)
    if department_name_raw.strip():
        return index.by_exact_name(department_name_raw)
    return None


def _parse_upload_code(raw: str) -> int:
    if not raw or not raw.isdigit() or len(raw) > 10:
        raise ValueError("ID клиента (КССС) — целое число до 10 знаков")
    code = int(raw)
    if code < 1 or code > 9999999998:
        raise ValueError("ID клиента (КССС) — целое число до 10 знаков")
    if code == SYSTEM_CLIENT_CODE:
        raise ValueError("Код 9999999999 зарезервирован")
    return code


async def export_clients_xlsx(db: AsyncSession) -> bytes:
    result = await db.execute(
        select(Client)
        .options(selectinload(Client.department))
        .where(Client.deleted_at.is_(None))
        .order_by(Client.name, Client.client_code)
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Клиенты"
    sheet.append(EXPORT_HEADERS)
    for client in result.scalars().all():
        department = client.department
        sheet.append(
            [
                client.client_code,
                client.name,
                client.inn,
                client.kpp or "",
                str(client.department_id) if client.department_id else "",
                department.name if department is not None else "",
                "Активен" if client.is_active else "Скрыт",
            ]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@dataclass
class _PreparedClient:
    client: Client
    department_name: str | None


@dataclass
class _ClientImportPlan:
    total: int
    pending: list[_PreparedClient]
    errors: list[ClientUploadError]


def _cell_at(row: tuple[object, ...], columns: dict[str, int], key: str) -> str:
    index = columns.get(key)
    if index is None or index >= len(row):
        return ""
    return _cell_text(row[index])


async def _plan_clients_xlsx(
    db: AsyncSession, content: bytes, filename: str
) -> _ClientImportPlan:
    if not filename.lower().endswith(".xlsx"):
        raise APIError(400, "INVALID_FILE", "Загрузите файл формата .xlsx")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise APIError(400, "INVALID_FILE", "Не удалось прочитать Excel") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise APIError(400, "INVALID_FILE", "Файл пуст")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise APIError(400, "INVALID_FILE", "Файл пуст")
        columns = _header_map(header)
        missing = {"client_code", "name", "inn"} - set(columns)
        if missing:
            raise APIError(400, "INVALID_FILE", "В файле нет обязательных колонок")
        departments = (
            await db.scalars(select(Department).where(Department.deleted_at.is_(None)))
        ).all()
        index = _DepartmentIndex(list(departments))
        existing_codes = set(
            (
                await db.scalars(
                    select(Client.client_code).where(Client.deleted_at.is_(None))
                )
            ).all()
        )
        seen_codes: set[int] = set()
        pending: list[_PreparedClient] = []
        errors: list[ClientUploadError] = []
        total = 0
        for row_number, row in enumerate(rows, start=2):
            if not any(_cell_text(value) for value in row):
                continue
            total += 1
            try:
                code = _parse_upload_code(_cell_at(row, columns, "client_code"))
                name = _validate_name(_cell_at(row, columns, "name"))
                inn = _validate_inn(_cell_at(row, columns, "inn"))
                kpp = _validate_kpp(_cell_at(row, columns, "kpp"))
                department = _resolve_department(
                    index,
                    _cell_at(row, columns, "department"),
                    _cell_at(row, columns, "department_name"),
                )
                if code in existing_codes or code in seen_codes:
                    raise ValueError(f"Клиент с кодом {code} уже есть")
                seen_codes.add(code)
                pending.append(
                    _PreparedClient(
                        client=Client(
                            client_code=code,
                            inn=inn,
                            name=name,
                            kpp=kpp,
                            department_id=department.id if department else None,
                            is_active=True,
                            is_system=False,
                        ),
                        department_name=department.name if department else None,
                    )
                )
            except (APIError, ValueError) as exc:
                message = exc.message if isinstance(exc, APIError) else str(exc)
                errors.append(ClientUploadError(row=row_number, message=message))
        if total == 0:
            raise APIError(400, "INVALID_FILE", "Файл пуст")
        return _ClientImportPlan(total=total, pending=pending, errors=errors)
    finally:
        workbook.close()


def _department_plans(
    pending: list[_PreparedClient],
) -> list[ClientUploadDepartmentPlan]:
    buckets: dict[tuple[UUID | None, str], int] = {}
    for item in pending:
        label = item.department_name or NO_DEPARTMENT
        key = (item.client.department_id, label)
        buckets[key] = buckets.get(key, 0) + 1
    plans = [
        ClientUploadDepartmentPlan(
            department_id=department_id,
            department_name=label,
            count=count,
        )
        for (department_id, label), count in buckets.items()
    ]
    plans.sort(key=lambda item: (-item.count, item.department_name))
    return plans


def _count_phrase(count: int, one: str, few: str, many: str) -> str:
    rest = abs(count) % 100
    last = rest % 10
    if 11 <= rest <= 14 or last == 0 or last >= 5:
        word = many
    elif last == 1:
        word = one
    else:
        word = few
    return f"{count} {word}"


def _preview_message(
    plans: list[ClientUploadDepartmentPlan], *, to_create: int, errors: int
) -> str:
    if to_create == 0:
        text = "Новые клиенты не будут созданы"
    else:
        listing = ", ".join(f"{item.department_name} — {item.count}" for item in plans)
        clients = _count_phrase(to_create, "клиент", "клиента", "клиентов")
        text = f"К загрузке {clients}: {listing}"
    if errors:
        rows = _count_phrase(errors, "строка", "строки", "строк")
        text += f". С ошибками {rows} — они не будут загружены"
    return text


async def preview_clients_xlsx(
    db: AsyncSession, content: bytes, filename: str
) -> ClientUploadPreview:
    plan = await _plan_clients_xlsx(db, content, filename)
    departments = _department_plans(plan.pending)
    without_department = sum(
        item.count for item in departments if item.department_id is None
    )
    return ClientUploadPreview(
        file_rows=plan.total,
        to_create=len(plan.pending),
        errors=len(plan.errors),
        without_department=without_department,
        departments=departments,
        message=_preview_message(
            departments, to_create=len(plan.pending), errors=len(plan.errors)
        ),
        error_details=plan.errors,
    )


async def upload_clients_xlsx(
    db: AsyncSession, content: bytes, filename: str
) -> ClientUploadResult:
    plan = await _plan_clients_xlsx(db, content, filename)
    for item in plan.pending:
        db.add(item.client)
    if plan.pending:
        await db.commit()
    message = f"Загружено {len(plan.pending)} из {plan.total}"
    return ClientUploadResult(
        created=len(plan.pending),
        errors=len(plan.errors),
        message=message,
        error_details=plan.errors,
    )


async def resolve_production_client(
    db: AsyncSession, client_code: int | None
) -> Client:
    if client_code is None:
        return await system_client(db)
    client = await client_by_code(db, client_code)
    if client is None:
        raise ValueError(f"Клиент с кодом {client_code} не найден")
    return client
