"""Two-file ERP balance upload: available stock and movements."""

from __future__ import annotations

import logging
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any, Literal
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import delete, select, tuple_
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import APIError
from app.models.available_balance import AvailableBalance
from app.models.available_balance_history import AvailableBalanceHistory
from app.models.object import Object
from app.models.product import Product
from app.models.sync_metadata import SyncMetadata
from app.schemas.logistics import (
    AvailablePreviewRow,
    BalancePreviewResult,
    BalanceUploadError,
    BalanceUploadResult,
    MovingPreviewRow,
)
from app.services.logistics_normative import (
    _is_blank,
    _parse_int_cell,
    _parse_quantity,
    _parse_warehouse_code,
)

logger = logging.getLogger(__name__)

FileKind = Literal["available", "moving"]
PREVIEW_LIMIT = 10
ERROR_LIMIT = 300
STORED_UNIT = "кг"
CENT = Decimal("0.01")
TON_TO_KG = Decimal("1000")

AVAILABLE_COLUMNS = (
    "КССС материала",
    "Завод",
    "Склад",
    "Запас",
    "Доступно из запаса",
    "Всего поступлений",
)
MOVING_COLUMNS = (
    "Склад-получатель",
    "№ продукта",
    "Открыто (не включено в ИП)",
    "Вес нетто",
    "Стат. движ.мат.",
    "Транспортный заказ",
    "ПродажнЕдиница",
)

KG_UNITS = {"КГ", "KG"}
PIECE_UNITS = {"ШТ", "PCS"}
TON_UNITS = {"Т", "T", "ТОНН", "ТОННА"}


@dataclass
class StockLine:
    warehouse_code: int
    product_code: int
    full_stock: Decimal
    available: Decimal
    on_way: Decimal
    no_ready: Decimal = Decimal("0.00")
    ready: Decimal = Decimal("0.00")
    collected: Decimal = Decimal("0.00")

    @property
    def plan(self) -> Decimal:
        return _quantize(
            self.available + self.no_ready + self.ready + self.collected + self.on_way
        )


@dataclass
class _AvailableRow:
    excel_row: int
    product_code: int
    erp_plant: int
    erp_warehouse: str
    full_stock: Decimal
    available: Decimal
    on_way: Decimal
    warehouse_code: int | None = None


@dataclass
class _MovingRow:
    excel_row: int
    product_code: int
    erp_warehouse: str
    status: str | None
    net_weight: Decimal
    opened: Decimal
    sales_unit: str | None
    transport_filled: bool
    warehouse_code: int | None = None


@dataclass
class _Prepared:
    available_preview: list[AvailablePreviewRow]
    moving_preview: list[MovingPreviewRow]
    available_rows: int
    moving_rows: int
    errors: list[BalanceUploadError] = field(default_factory=list)
    warnings: list[BalanceUploadError] = field(default_factory=list)
    lines: dict[tuple[int, int], StockLine] = field(default_factory=dict)
    replace_warehouses: set[int] = field(default_factory=set)

    @property
    def can_upload(self) -> bool:
        return not self.errors and bool(self.lines)


def _quantize(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def _display(value: Any) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _load_rows(content: bytes, filename: str) -> list[tuple[Any, ...]]:
    lowered = (filename or "").lower()
    if not lowered.endswith((".xlsx", ".xls")):
        raise APIError(400, "INVALID_FILE", "Загрузите файл .xlsx или .xls")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise APIError(
            400,
            "INVALID_FILE",
            "Не удалось прочитать Excel. Используйте формат .xlsx",
        ) from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise APIError(400, "INVALID_FILE", "Файл пуст")
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows:
        raise APIError(400, "INVALID_FILE", "Файл пуст")
    return rows


def _header_map(row: tuple[Any, ...]) -> tuple[dict[str, int], list[str]]:
    found: dict[str, int] = {}
    duplicates: list[str] = []
    for index, cell in enumerate(row):
        if _is_blank(cell):
            continue
        name = str(cell).strip()
        if name in found:
            duplicates.append(name)
        found[name] = index
    return found, duplicates


def _value(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    index = headers.get(name)
    if index is None or index >= len(row):
        return None
    return row[index]


def _blank_row(row: tuple[Any, ...] | None) -> bool:
    if row is None:
        return True
    return all(_is_blank(cell) for cell in row)


def _non_negative(value: Any, title: str) -> Decimal:
    amount = _parse_quantity(value, f"«{title}»")
    if amount < 0:
        raise ValueError(f"«{title}» должно быть числом не меньше 0")
    return amount


def _sales_unit(value: Any) -> str:
    if _is_blank(value):
        raise ValueError("Укажите «ПродажнЕдиница»")
    text = str(value).strip().upper().replace(" ", "").rstrip(".")
    if text in KG_UNITS:
        return "KG"
    if text in PIECE_UNITS:
        return "PIECE"
    if text in TON_UNITS:
        return "TON"
    raise ValueError(
        f"ПродажнЕдиница «{value}» не переводится в кг (ожидаются КГ, ШТ или т)"
    )


def _movement_status(value: Any) -> str | None:
    if _is_blank(value):
        return None
    return str(value).strip().upper().replace("А", "A")


def _to_kg(
    amount: Decimal,
    unit: str,
    weight_kg: Decimal | None,
    product_code: int,
) -> Decimal:
    if unit == "KG":
        return amount
    if unit == "TON":
        return _quantize(amount * TON_TO_KG)
    if weight_kg is None or weight_kg <= 0:
        raise ValueError(
            f"У артикула {product_code} не задан weight_kg — нельзя перевести ШТ в кг"
        )
    return _quantize(amount * weight_kg)


def _error(kind: FileKind, row: int, message: str) -> BalanceUploadError:
    return BalanceUploadError(row=row, message=message, file=kind)


def _column_errors(
    kind: FileKind,
    headers: dict[str, int],
    duplicates: list[str],
    required: tuple[str, ...],
) -> list[BalanceUploadError]:
    errors: list[BalanceUploadError] = []
    for name in required:
        if name not in headers:
            errors.append(_error(kind, 1, f"Нет колонки «{name}»"))
    for name in duplicates:
        errors.append(_error(kind, 1, f"Колонка «{name}» повторяется"))
    return errors


def _available_preview(
    rows: list[tuple[Any, ...]], headers: dict[str, int]
) -> tuple[list[AvailablePreviewRow], int]:
    preview: list[AvailablePreviewRow] = []
    count = 0
    for excel_row, row in enumerate(rows[1:], start=2):
        if _blank_row(row):
            continue
        count += 1
        if len(preview) < PREVIEW_LIMIT:
            preview.append(
                AvailablePreviewRow(
                    row=excel_row,
                    product_code=_display(_value(row, headers, "КССС материала")),
                    plant=_display(_value(row, headers, "Завод")),
                    warehouse=_display(_value(row, headers, "Склад")),
                    full_stock=_display(_value(row, headers, "Запас")),
                    available=_display(_value(row, headers, "Доступно из запаса")),
                    on_way=_display(_value(row, headers, "Всего поступлений")),
                )
            )
    return preview, count


def _moving_preview(
    rows: list[tuple[Any, ...]], headers: dict[str, int]
) -> tuple[list[MovingPreviewRow], int]:
    preview: list[MovingPreviewRow] = []
    count = 0
    for excel_row, row in enumerate(rows[1:], start=2):
        if _blank_row(row):
            continue
        count += 1
        if len(preview) < PREVIEW_LIMIT:
            preview.append(
                MovingPreviewRow(
                    row=excel_row,
                    product_code=_display(_value(row, headers, "№ продукта")),
                    warehouse=_display(_value(row, headers, "Склад-получатель")),
                    status=_display(_value(row, headers, "Стат. движ.мат.")),
                    net_weight=_display(_value(row, headers, "Вес нетто")),
                    opened=_display(_value(row, headers, "Открыто (не включено в ИП)")),
                    sales_unit=_display(_value(row, headers, "ПродажнЕдиница")),
                    transport_order=_display(
                        _value(row, headers, "Транспортный заказ")
                    ),
                )
            )
    return preview, count


def _parse_available(
    rows: list[tuple[Any, ...]], headers: dict[str, int]
) -> tuple[list[_AvailableRow], list[BalanceUploadError]]:
    parsed: list[_AvailableRow] = []
    errors: list[BalanceUploadError] = []
    for excel_row, row in enumerate(rows[1:], start=2):
        if _blank_row(row):
            continue
        try:
            parsed.append(
                _AvailableRow(
                    excel_row=excel_row,
                    product_code=_parse_int_cell(
                        _value(row, headers, "КССС материала"),
                        "«КССС материала»",
                    ),
                    erp_plant=_parse_int_cell(
                        _value(row, headers, "Завод"),
                        "«Завод»",
                    ),
                    erp_warehouse=_parse_warehouse_code(_value(row, headers, "Склад")),
                    full_stock=_non_negative(_value(row, headers, "Запас"), "Запас"),
                    available=_parse_quantity(
                        _value(row, headers, "Доступно из запаса"),
                        "«Доступно из запаса»",
                    ),
                    on_way=_non_negative(
                        _value(row, headers, "Всего поступлений"),
                        "Всего поступлений",
                    ),
                )
            )
        except ValueError as exc:
            errors.append(_error("available", excel_row, str(exc)))
    return parsed, errors


def _parse_moving(
    rows: list[tuple[Any, ...]], headers: dict[str, int]
) -> tuple[list[_MovingRow], list[BalanceUploadError]]:
    parsed: list[_MovingRow] = []
    errors: list[BalanceUploadError] = []
    for excel_row, row in enumerate(rows[1:], start=2):
        if _blank_row(row):
            continue
        try:
            status = _movement_status(_value(row, headers, "Стат. движ.мат."))
            sales_unit = None
            if status is None:
                sales_unit = _sales_unit(_value(row, headers, "ПродажнЕдиница"))
            parsed.append(
                _MovingRow(
                    excel_row=excel_row,
                    product_code=_parse_int_cell(
                        _value(row, headers, "№ продукта"),
                        "«№ продукта»",
                    ),
                    erp_warehouse=_parse_warehouse_code(
                        _value(row, headers, "Склад-получатель")
                    ),
                    status=status,
                    net_weight=_non_negative(
                        _value(row, headers, "Вес нетто"),
                        "Вес нетто",
                    ),
                    opened=_non_negative(
                        _value(row, headers, "Открыто (не включено в ИП)"),
                        "Открыто (не включено в ИП)",
                    ),
                    sales_unit=sales_unit,
                    transport_filled=not _is_blank(
                        _value(row, headers, "Транспортный заказ")
                    ),
                )
            )
        except ValueError as exc:
            errors.append(_error("moving", excel_row, str(exc)))
    return parsed, errors


def _duplicate_errors(rows: list[_AvailableRow]) -> list[BalanceUploadError]:
    first: dict[tuple[int, int, str], int] = {}
    errors: list[BalanceUploadError] = []
    for row in rows:
        key = (row.product_code, row.erp_plant, row.erp_warehouse)
        if key not in first:
            first[key] = row.excel_row
            continue
        errors.append(
            _error(
                "available",
                row.excel_row,
                (
                    f"Повтор пары артикул {row.product_code} "
                    f"и склад {row.erp_plant}/{row.erp_warehouse} "
                    f"(первая строка {first[key]})"
                ),
            )
        )
    return errors


def _prefer_object(items: list[Object]) -> Object:
    return sorted(items, key=lambda item: (not item.is_active, item.code))[0]


async def _load_products(db: AsyncSession, codes: set[int]) -> dict[int, Decimal]:
    if not codes:
        return {}
    result = await db.execute(
        select(Product.code, Product.weight_kg).where(
            Product.code.in_(codes),
            Product.deleted_at.is_(None),
        )
    )
    return {code: weight for code, weight in result.all()}


async def _load_warehouses(
    db: AsyncSession, erp_codes: set[str]
) -> tuple[dict[tuple[int, str], Object], dict[str, Object]]:
    if not erp_codes:
        return {}, {}
    result = await db.execute(
        select(Object).where(
            Object.type == "warehouse",
            Object.deleted_at.is_(None),
            Object.erp_warehouse_code.in_(erp_codes),
        )
    )
    by_pair: dict[tuple[int, str], list[Object]] = defaultdict(list)
    by_code: dict[str, list[Object]] = defaultdict(list)
    for item in result.scalars().all():
        warehouse_code = item.erp_warehouse_code or ""
        by_code[warehouse_code].append(item)
        if item.erp_plant_code is not None:
            by_pair[(item.erp_plant_code, warehouse_code)].append(item)
    return (
        {key: _prefer_object(items) for key, items in by_pair.items()},
        {key: _prefer_object(items) for key, items in by_code.items()},
    )


def _message(prepared: _Prepared) -> str:
    if prepared.errors:
        shown = min(len(prepared.errors), ERROR_LIMIT)
        if len(prepared.errors) > shown:
            return (
                f"Проверка не пройдена: {len(prepared.errors)} ошибок. "
                f"Показаны первые {shown}"
            )
        return f"Проверка не пройдена: {len(prepared.errors)} ошибок"
    skipped = len(prepared.warnings)
    if not prepared.lines:
        if skipped:
            return f"Нет строк с известным артикулом. Пропущено {skipped}"
        return "Нет строк для загрузки"
    if skipped:
        return (
            f"Можно загрузить {len(prepared.lines)} позиций. "
            f"Пропущено {skipped}: артикул не найден"
        )
    return f"Можно загрузить {len(prepared.lines)} позиций"


def _public_errors(errors: list[BalanceUploadError]) -> list[BalanceUploadError]:
    ordered = sorted(errors, key=lambda item: (item.file or "", item.row, item.message))
    return ordered[:ERROR_LIMIT]


async def _prepare(
    db: AsyncSession,
    available_content: bytes,
    available_name: str,
    moving_content: bytes,
    moving_name: str,
) -> _Prepared:
    available_sheet = _load_rows(available_content, available_name)
    moving_sheet = _load_rows(moving_content, moving_name)
    available_headers, available_dups = _header_map(available_sheet[0])
    moving_headers, moving_dups = _header_map(moving_sheet[0])
    available_preview, available_count = _available_preview(
        available_sheet, available_headers
    )
    moving_preview, moving_count = _moving_preview(moving_sheet, moving_headers)
    errors = _column_errors(
        "available", available_headers, available_dups, AVAILABLE_COLUMNS
    ) + _column_errors("moving", moving_headers, moving_dups, MOVING_COLUMNS)
    warnings: list[BalanceUploadError] = []

    available_rows: list[_AvailableRow] = []
    moving_rows: list[_MovingRow] = []
    available_structure_ok = not any(item.file == "available" for item in errors)
    moving_structure_ok = not any(item.file == "moving" for item in errors)
    if available_structure_ok:
        available_rows, row_errors = _parse_available(
            available_sheet, available_headers
        )
        errors.extend(row_errors)
        errors.extend(_duplicate_errors(available_rows))
        if available_count == 0:
            errors.append(
                _error("available", 1, "В файле «Доступные остатки» нет строк")
            )
    if moving_structure_ok:
        moving_rows, row_errors = _parse_moving(moving_sheet, moving_headers)
        errors.extend(row_errors)

    product_codes = {row.product_code for row in available_rows}
    product_codes.update(row.product_code for row in moving_rows)
    erp_codes = {row.erp_warehouse for row in available_rows}
    erp_codes.update(row.erp_warehouse for row in moving_rows)
    products = await _load_products(db, product_codes)
    by_pair, by_code = await _load_warehouses(db, erp_codes)

    lines: dict[tuple[int, int], StockLine] = {}
    for row in available_rows:
        if row.product_code not in products:
            warnings.append(
                _error(
                    "available",
                    row.excel_row,
                    f"Артикул {row.product_code} не найден",
                )
            )
        warehouse = by_pair.get((row.erp_plant, row.erp_warehouse))
        if warehouse is None:
            errors.append(
                _error(
                    "available",
                    row.excel_row,
                    f"Склад ERP {row.erp_plant}/{row.erp_warehouse} не найден",
                )
            )
            continue
        row.warehouse_code = warehouse.code
        if row.product_code not in products:
            continue
        lines[(warehouse.code, row.product_code)] = StockLine(
            warehouse_code=warehouse.code,
            product_code=row.product_code,
            full_stock=row.full_stock,
            available=row.available,
            on_way=row.on_way,
        )

    replace_warehouses = {line.warehouse_code for line in lines.values()}
    if not available_structure_ok:
        return _Prepared(
            available_preview=available_preview,
            moving_preview=moving_preview,
            available_rows=available_count,
            moving_rows=moving_count,
            errors=errors,
            warnings=warnings,
            lines={},
        )

    buckets: dict[tuple[int, int], dict[str, Decimal]] = defaultdict(
        lambda: {
            "no_ready": Decimal("0.00"),
            "ready": Decimal("0.00"),
            "collected": Decimal("0.00"),
        }
    )
    for row in moving_rows:
        if row.product_code not in products:
            warnings.append(
                _error("moving", row.excel_row, f"Артикул {row.product_code} не найден")
            )
        warehouse = by_code.get(row.erp_warehouse)
        if warehouse is None:
            errors.append(
                _error(
                    "moving",
                    row.excel_row,
                    f"Склад-получатель {row.erp_warehouse} не найден",
                )
            )
            continue
        row.warehouse_code = warehouse.code
        if row.product_code not in products:
            continue
        key = (warehouse.code, row.product_code)
        if key not in lines:
            lines[key] = StockLine(
                warehouse_code=warehouse.code,
                product_code=row.product_code,
                full_stock=Decimal("0.00"),
                available=Decimal("0.00"),
                on_way=Decimal("0.00"),
            )
        try:
            if row.status is None:
                if row.opened == 0:
                    addition = Decimal("0.00")
                else:
                    addition = _to_kg(
                        row.opened,
                        row.sales_unit or "KG",
                        products[row.product_code],
                        row.product_code,
                    )
                buckets[key]["no_ready"] += addition
            elif row.status == "A" and not row.transport_filled:
                buckets[key]["ready"] += row.net_weight
            elif row.status == "A" and row.transport_filled:
                buckets[key]["collected"] += row.net_weight
        except ValueError as exc:
            errors.append(_error("moving", row.excel_row, str(exc)))

    for key, bucket in buckets.items():
        line = lines.get(key)
        if line is None:
            continue
        line.no_ready = _quantize(bucket["no_ready"])
        line.ready = _quantize(bucket["ready"])
        line.collected = _quantize(bucket["collected"])

    return _Prepared(
        available_preview=available_preview,
        moving_preview=moving_preview,
        available_rows=available_count,
        moving_rows=moving_count,
        errors=errors,
        warnings=warnings,
        lines=lines,
        replace_warehouses=replace_warehouses,
    )


def _result(prepared: _Prepared) -> BalancePreviewResult:
    return BalancePreviewResult(
        available_preview=prepared.available_preview,
        moving_preview=prepared.moving_preview,
        available_rows=prepared.available_rows,
        moving_rows=prepared.moving_rows,
        errors=len(prepared.errors),
        warnings=len(prepared.warnings),
        message=_message(prepared),
        can_upload=prepared.can_upload,
        error_details=_public_errors(prepared.errors),
        warning_details=_public_errors(prepared.warnings),
    )


async def preview_balances(
    db: AsyncSession,
    available_content: bytes,
    available_name: str,
    moving_content: bytes,
    moving_name: str,
) -> BalancePreviewResult:
    prepared = await _prepare(
        db,
        available_content,
        available_name,
        moving_content,
        moving_name,
    )
    return _result(prepared)


async def _touch_sync_metadata(
    db: AsyncSession, *, user_id: UUID, synced_at: datetime
) -> None:
    row = await db.get(SyncMetadata, 1)
    if row is None:
        db.add(
            SyncMetadata(
                id=1,
                last_balances_sync_at=synced_at,
                last_balances_sync_by=user_id,
                updated_at=synced_at,
            )
        )
        return
    row.last_balances_sync_at = synced_at
    row.last_balances_sync_by = user_id
    row.updated_at = synced_at


async def upload_balances(
    db: AsyncSession,
    available_content: bytes,
    available_name: str,
    moving_content: bytes,
    moving_name: str,
    *,
    user_id: UUID,
) -> BalanceUploadResult:
    prepared = await _prepare(
        db,
        available_content,
        available_name,
        moving_content,
        moving_name,
    )
    if not prepared.can_upload:
        raise APIError(400, "VALIDATION_ERROR", _message(prepared))

    warehouse_codes = {line.warehouse_code for line in prepared.lines.values()}
    replace_warehouses = prepared.replace_warehouses
    extra_keys = [
        key
        for key, line in prepared.lines.items()
        if line.warehouse_code not in replace_warehouses
    ]
    existing = await db.execute(
        select(AvailableBalance.warehouse_code, AvailableBalance.product_code).where(
            AvailableBalance.warehouse_code.in_(warehouse_codes)
        )
    )
    existing_keys = {(row[0], row[1]) for row in existing.all()}
    now = datetime.now(UTC)
    snapshot_date = now.date()
    if replace_warehouses:
        await db.execute(
            delete(AvailableBalance).where(
                AvailableBalance.warehouse_code.in_(replace_warehouses)
            )
        )
        await db.execute(
            delete(AvailableBalanceHistory).where(
                AvailableBalanceHistory.snapshot_date == snapshot_date,
                AvailableBalanceHistory.warehouse_code.in_(replace_warehouses),
            )
        )
    if extra_keys:
        balance_key = tuple_(
            AvailableBalance.warehouse_code,
            AvailableBalance.product_code,
        )
        history_key = tuple_(
            AvailableBalanceHistory.warehouse_code,
            AvailableBalanceHistory.product_code,
        )
        await db.execute(delete(AvailableBalance).where(balance_key.in_(extra_keys)))
        await db.execute(
            delete(AvailableBalanceHistory).where(
                AvailableBalanceHistory.snapshot_date == snapshot_date,
                history_key.in_(extra_keys),
            )
        )
    await db.flush()

    created = 0
    updated = 0
    for key, line in prepared.lines.items():
        db.add(
            AvailableBalance(
                warehouse_code=line.warehouse_code,
                product_code=line.product_code,
                available=line.available,
                plan=line.plan,
                full_stock=line.full_stock,
                on_way=line.on_way,
                no_ready=line.no_ready,
                ready=line.ready,
                collected=line.collected,
                unit=STORED_UNIT,
                last_sync_at=now,
                source="excel_upload",
            )
        )
        db.add(
            AvailableBalanceHistory(
                snapshot_date=snapshot_date,
                snapshot_at=now,
                warehouse_code=line.warehouse_code,
                product_code=line.product_code,
                available=line.available,
                plan=line.plan,
                full_stock=line.full_stock,
                unit=STORED_UNIT,
                source="excel_upload",
                uploaded_by=user_id,
            )
        )
        if key in existing_keys:
            updated += 1
        else:
            created += 1

    await _touch_sync_metadata(db, user_id=user_id, synced_at=now)
    await db.commit()
    from app.services.dashboard import refresh_mv_normative_daily

    await refresh_mv_normative_daily(db)
    await db.commit()
    loaded = created + updated
    skipped = len(prepared.warnings)
    message = f"Загружено {loaded}"
    if skipped:
        message = f"{message}. Пропущено {skipped}: артикул не найден"
    logger.info(
        "Balance upload: positions=%s created=%s updated=%s skipped=%s warehouses=%s",
        loaded,
        created,
        updated,
        skipped,
        sorted(warehouse_codes),
    )
    return BalanceUploadResult(
        uploaded=loaded,
        created=created,
        updated=updated,
        errors=0,
        warnings=skipped,
        message=message,
        error_details=[],
        warning_details=_public_errors(prepared.warnings),
    )
