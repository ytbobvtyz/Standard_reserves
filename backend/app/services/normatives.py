from collections import defaultdict
from datetime import UTC, date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Literal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import String, and_, cast, exists, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import APIError
from app.core.pagination import paginate
from app.models.normative import Normative
from app.models.object import Object
from app.models.product import Product
from app.models.production_request import ProductionRequest, ProductionRequestItem
from app.models.request import Request
from app.models.user import User
from app.schemas.common import PaginationMeta
from app.schemas.normative import (
    NormativeCalculateData,
    NormativeListItem,
    NormativeOnDateDetail,
    NormativeOnDateItem,
)
from app.services.coefficients import (
    calculate_requirement,
    category_factor,
    distance_factor,
)
from app.services.params import load_coefficient_set

DEFAULT_UNIT = "шт"
EXPORT_HEADERS = [
    "Артикул",
    "Название",
    "Склад",
    "Запрос",
    "Автор",
    "Количество",
    "Ед.",
    "Срок действия",
    "Клиент",
    "Подразделение",
]
EXPORT_COLUMN_WIDTHS = [14, 36, 28, 14, 24, 14, 8, 16, 28, 28]

CategoryFilter = Literal["A", "B", "C"]


def _end_of_day(value: date) -> datetime:
    return datetime.combine(value, time.max, tzinfo=UTC)


def _product_search_condition(search: str | None):
    if not search or not search.strip():
        return None
    term = f"%{search.strip()}%"
    return exists(
        select(1).where(
            Product.code == Normative.product_code,
            or_(
                cast(Product.code, String).ilike(term),
                Product.name.ilike(term),
            ),
        )
    )


def _department_fields(normative: Normative) -> tuple[UUID | None, str | None]:
    request = normative.request
    if request is not None:
        department = request.department
        if department is not None:
            return request.department_id, department.name
        initiator = request.initiator
        return request.department_id, initiator.department if initiator else None

    production_item = normative.production_request_item
    if production_item is None:
        return None, None
    uploader = production_item.production_request.uploader
    department = uploader.assigned_department
    if department is not None:
        return uploader.department_id, department.name
    return uploader.department_id, uploader.department


def _production_valid_from_condition(on_date: date):
    return exists(
        select(1)
        .select_from(ProductionRequestItem)
        .join(
            ProductionRequest,
            ProductionRequest.id == ProductionRequestItem.production_request_id,
        )
        .where(
            ProductionRequestItem.id == Normative.production_request_item_id,
            ProductionRequest.deleted_at.is_(None),
            ProductionRequest.status == "active",
            ProductionRequest.valid_from <= on_date,
        )
    )


def _valid_on_condition(on_date: date):
    return or_(
        and_(
            Normative.production_request_item_id.is_(None),
            Normative.created_at <= _end_of_day(on_date),
        ),
        _production_valid_from_condition(on_date),
    )


def _department_condition(department_id: UUID):
    regular_request = exists(
        select(1).where(
            Request.id == Normative.request_id,
            Request.department_id == department_id,
        )
    )
    production_upload = exists(
        select(1)
        .select_from(ProductionRequestItem)
        .join(
            ProductionRequest,
            ProductionRequest.id == ProductionRequestItem.production_request_id,
        )
        .join(User, User.id == ProductionRequest.uploaded_by)
        .where(
            ProductionRequestItem.id == Normative.production_request_item_id,
            User.department_id == department_id,
        )
    )
    return or_(regular_request, production_upload)


def _load_options():
    return (
        selectinload(Normative.product),
        selectinload(Normative.warehouse),
        selectinload(Normative.request).selectinload(Request.department),
        selectinload(Normative.request).selectinload(Request.initiator),
        selectinload(Normative.production_request_item)
        .selectinload(ProductionRequestItem.production_request)
        .selectinload(ProductionRequest.uploader)
        .selectinload(User.assigned_department),
    )


def _author_name(normative: Normative) -> str | None:
    if normative.request and normative.request.initiator:
        return normative.request.initiator.full_name
    if (
        normative.production_request_item
        and normative.production_request_item.production_request
        and normative.production_request_item.production_request.uploader
    ):
        return normative.production_request_item.production_request.uploader.full_name
    return None


def to_list_item(normative: Normative) -> NormativeListItem:
    department_id, department_name = _department_fields(normative)
    return NormativeListItem(
        id=normative.id,
        request_id=normative.request_id,
        author_name=_author_name(normative),
        product_code=normative.product_code,
        product_name=normative.product.name if normative.product else "",
        category=normative.category.strip(),
        warehouse_code=normative.warehouse_code,
        warehouse_name=normative.warehouse.name if normative.warehouse else "",
        quantity=normative.quantity,
        unit=normative.unit,
        client_name=normative.client_name,
        department_id=department_id,
        department_name=department_name,
        expiry_date=normative.expiry_date,
        created_at=normative.created_at,
    )


async def list_current_normatives(
    db: AsyncSession,
    *,
    warehouse_code: int | None,
    product_code: int | None,
    client_name: str | None,
    category: CategoryFilter | None,
    search: str | None,
    department_id: UUID | None,
    page: int,
    limit: int,
) -> tuple[list[NormativeListItem], PaginationMeta]:
    conditions = [
        Normative.deleted_at.is_(None),
        Normative.expiry_date >= date.today(),
        _valid_on_condition(date.today()),
    ]
    if warehouse_code is not None:
        conditions.append(Normative.warehouse_code == warehouse_code)
    if product_code is not None:
        conditions.append(Normative.product_code == product_code)
    if client_name:
        conditions.append(Normative.client_name.ilike(f"%{client_name.strip()}%"))
    if category:
        conditions.append(Normative.category == category)
    if department_id is not None:
        conditions.append(_department_condition(department_id))
    search_condition = _product_search_condition(search)
    if search_condition is not None:
        conditions.append(search_condition)

    total = await db.scalar(
        select(func.count()).select_from(Normative).where(*conditions)
    )
    result = await db.execute(
        paginate(
            select(Normative)
            .options(*_load_options())
            .where(*conditions)
            .order_by(
                Normative.warehouse_code,
                Normative.product_code,
                Normative.created_at.desc(),
            ),
            page,
            limit,
        )
    )
    items = [to_list_item(row) for row in result.scalars().unique().all()]
    return items, PaginationMeta(page=page, limit=limit, total=total or 0)


async def list_normatives_on_date(
    db: AsyncSession,
    *,
    on_date: date,
    warehouse_code: int | None,
    product_code: int | None,
    search: str | None = None,
    department_id: UUID | None = None,
) -> list[NormativeOnDateItem]:
    conditions = [
        Normative.deleted_at.is_(None),
        _valid_on_condition(on_date),
        Normative.expiry_date >= on_date,
    ]
    if warehouse_code is not None:
        conditions.append(Normative.warehouse_code == warehouse_code)
    if product_code is not None:
        conditions.append(Normative.product_code == product_code)
    if department_id is not None:
        conditions.append(_department_condition(department_id))
    search_condition = _product_search_condition(search)
    if search_condition is not None:
        conditions.append(search_condition)

    result = await db.execute(
        select(Normative)
        .options(*_load_options())
        .where(*conditions)
        .order_by(
            Normative.warehouse_code,
            Normative.product_code,
            Normative.created_at.desc(),
        )
    )
    grouped: dict[tuple[int, int, str], list[Normative]] = defaultdict(list)
    for normative in result.scalars().unique().all():
        key = (normative.warehouse_code, normative.product_code, normative.unit)
        grouped[key].append(normative)

    items: list[NormativeOnDateItem] = []
    for (_warehouse, _product, unit), rows in grouped.items():
        seen_clients: set[str] = set()
        latest_rows: list[Normative] = []
        for row in rows:
            client_key = row.client_name.strip().lower()
            if client_key not in seen_clients:
                seen_clients.add(client_key)
                latest_rows.append(row)

        first = latest_rows[0]
        items.append(
            NormativeOnDateItem(
                product_code=first.product_code,
                product_name=first.product.name if first.product else "",
                warehouse_code=first.warehouse_code,
                warehouse_name=first.warehouse.name if first.warehouse else "",
                total_quantity=sum((row.quantity for row in latest_rows), Decimal("0")),
                unit=unit,
                category=first.category.strip(),
                details=[
                    NormativeOnDateDetail(
                        client_name=row.client_name,
                        quantity=row.quantity,
                        expiry_date=row.expiry_date,
                        department_id=_department_fields(row)[0],
                        department_name=_department_fields(row)[1],
                        request_id=row.request_id,
                        author_name=_author_name(row),
                    )
                    for row in latest_rows
                ],
            )
        )
    items.sort(key=lambda item: (item.warehouse_code, item.product_code))
    return items


def _request_label(request_id: UUID | None) -> str:
    if request_id is None:
        return ""
    return f"№{str(request_id)[:8]}"


def build_export_xlsx(
    items: list[NormativeOnDateItem],
    *,
    category: CategoryFilter | None = None,
    client_name: str | None = None,
) -> bytes:
    client_filter = client_name.strip().lower() if client_name and client_name.strip() else None
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Нормативы"
    sheet.append(EXPORT_HEADERS)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADERS))}1"
    for index, width in enumerate(EXPORT_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 22

    for item in items:
        if category and (item.category or "").strip().upper() != category:
            continue
        for detail in item.details:
            if client_filter and client_filter not in detail.client_name.lower():
                continue
            sheet.append(
                [
                    item.product_code,
                    item.product_name,
                    item.warehouse_name,
                    _request_label(detail.request_id),
                    detail.author_name or "",
                    float(detail.quantity),
                    item.unit,
                    detail.expiry_date,
                    detail.client_name,
                    detail.department_name or "",
                ]
            )
            row_idx = sheet.max_row
            for col in range(1, len(EXPORT_HEADERS) + 1):
                cell = sheet.cell(row=row_idx, column=col)
                cell.border = thin
                cell.alignment = Alignment(vertical="center", wrap_text=col in {2, 3, 5, 9, 10})
            sheet.cell(row=row_idx, column=1).number_format = "0"
            sheet.cell(row=row_idx, column=6).number_format = "#,##0.00"
            sheet.cell(row=row_idx, column=8).number_format = "DD.MM.YYYY"

    if sheet.max_row > 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADERS))}{sheet.max_row}"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def export_normatives_xlsx(
    db: AsyncSession,
    *,
    on_date: date,
    warehouse_code: int | None = None,
    product_code: int | None = None,
    search: str | None = None,
    department_id: UUID | None = None,
    category: CategoryFilter | None = None,
    client_name: str | None = None,
) -> bytes:
    items = await list_normatives_on_date(
        db,
        on_date=on_date,
        warehouse_code=warehouse_code,
        product_code=product_code,
        search=search,
        department_id=department_id,
    )
    return build_export_xlsx(items, category=category, client_name=client_name)


async def calculate_normative(
    db: AsyncSession,
    *,
    product_code: int,
    warehouse_code: int,
) -> NormativeCalculateData:
    product = await db.scalar(
        select(Product).where(
            Product.code == product_code,
            Product.deleted_at.is_(None),
        )
    )
    if product is None:
        raise APIError(404, "NOT_FOUND", "Продукт не найден")

    warehouse = await db.scalar(
        select(Object).where(
            Object.code == warehouse_code,
            Object.deleted_at.is_(None),
        )
    )
    if warehouse is None:
        raise APIError(404, "NOT_FOUND", "Склад не найден")

    category = product.category.strip()
    coeffs = await load_coefficient_set(db)
    cat_factor = category_factor(category, coeffs)
    dist_factor = distance_factor(bool(warehouse.long_distance), coeffs)
    monthly = product.monthly_consumption
    calculated = None
    if monthly is not None:
        calculated = calculate_requirement(
            monthly, category, bool(warehouse.long_distance), coeffs
        )

    return NormativeCalculateData(
        product_code=product.code,
        product_name=product.name,
        category=category,
        warehouse_code=warehouse.code,
        warehouse_name=warehouse.name,
        monthly_consumption=monthly,
        distance_factor=dist_factor,
        category_factor=cat_factor,
        calculated_normative=calculated,
        unit=DEFAULT_UNIT,
    )
