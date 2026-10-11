from datetime import date
from io import BytesIO
from uuid import uuid4

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.database import AsyncSessionLocal
from app.models.client import SYSTEM_CLIENT_ID
from app.models.department import Department
from app.models.event import Event
from app.models.normative import Normative
from app.models.request import Request
from tests.conftest import (
    AuthUser,
    auth_header,
    delete_request,
    ensure_client,
    login_token,
)
from tests.test_approvals import _create_submitted
from tests.test_requests import valid_expiry_date


def _clients_xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["ID клиента (КССС)", "Наименование", "ИНН", "КПП", "ID подразделения"]
    )
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def test_create_client_inherits_department_and_rejects_duplicate(
    client: AsyncClient, test_user: AuthUser
) -> None:
    token = await login_token(client, test_user)
    payload = {
        "client_code": 7701000001,
        "inn": "7701000001",
        "name": "ООО Наследник",
    }
    created = await client.post(
        "/api/v1/references/clients",
        headers=auth_header(token),
        json=payload,
    )
    assert created.status_code == 201, created.text
    body = created.json()["data"]
    assert body["department_id"] is not None
    assert body["name"] == "ООО Наследник"
    duplicate = await client.post(
        "/api/v1/references/clients",
        headers=auth_header(token),
        json={**payload, "inn": "7701000002", "name": "Другое имя"},
    )
    assert duplicate.status_code == 409
    assert duplicate.json()["error"]["code"] == "CONFLICT"


async def test_search_by_code_inn_and_name(
    client: AsyncClient, test_user: AuthUser
) -> None:
    token = await login_token(client, test_user)
    await client.post(
        "/api/v1/references/clients",
        headers=auth_header(token),
        json={
            "client_code": 7702000002,
            "inn": "770200000212",
            "name": "Поиск Ромашка",
        },
    )
    headers = auth_header(token)
    by_code = await client.get(
        "/api/v1/references/clients",
        headers=headers,
        params={"q": "7702000002"},
    )
    by_inn = await client.get(
        "/api/v1/references/clients",
        headers=headers,
        params={"q": "770200000212"},
    )
    by_name = await client.get(
        "/api/v1/references/clients",
        headers=headers,
        params={"q": "ромаш"},
    )
    assert by_code.status_code == 200
    assert by_inn.status_code == 200
    assert by_name.status_code == 200
    assert by_code.json()["data"][0]["client_code"] == 7702000002
    assert by_inn.json()["data"][0]["inn"] == "770200000212"
    assert any(item["name"] == "Поиск Ромашка" for item in by_name.json()["data"])


async def test_system_client_cannot_be_changed(
    client: AsyncClient, logistics_user: AuthUser
) -> None:
    token = await login_token(client, logistics_user)
    response = await client.put(
        f"/api/v1/references/clients/{SYSTEM_CLIENT_ID}",
        headers=auth_header(token),
        json={"name": "Чужое имя", "is_active": True, "department_id": None},
    )
    assert response.status_code == 400


async def test_excel_department_by_name(
    client: AsyncClient, logistics_user: AuthUser
) -> None:
    token = await login_token(client, logistics_user)
    department_name = f"Sales Office {uuid4().hex[:6]}"
    async with AsyncSessionLocal() as session:
        department = Department(name=department_name, is_active=True)
        session.add(department)
        await session.commit()
        department_id = department.id
    content = _clients_xlsx(
        [[7703000003, "Excel Клиент", "7703000003", "", department_name.lower()]]
    )
    response = await client.post(
        "/api/v1/references/clients/upload",
        headers=auth_header(token),
        files={
            "file": (
                "clients.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["data"]["created"] == 1
    listed = await client.get(
        "/api/v1/references/clients",
        headers=auth_header(token),
        params={"q": "7703000003", "include_inactive": True},
    )
    match = next(
        item for item in listed.json()["data"] if item["client_code"] == 7703000003
    )
    assert match["department_id"] == str(department_id)

    template = await client.get(
        "/api/v1/references/clients/template",
        headers=auth_header(token),
    )
    book = load_workbook(BytesIO(template.content))
    assert book.active["E1"].comment is not None
    assert "Коммерческий отдел" in book.active["E1"].comment.text


async def test_name_replacement_writes_event_and_same_client_does_not(
    client: AsyncClient,
    test_user: AuthUser,
    pp_user: AuthUser,
    economist_user: AuthUser,
    catalog: dict[str, int],
) -> None:
    commercial_token = await login_token(client, test_user)
    pp_token = await login_token(client, pp_user)
    economist_token = await login_token(client, economist_user)
    name = "Ромашка Замещение"
    client_id = await ensure_client(name)
    old_id = uuid4()
    async with AsyncSessionLocal() as session:
        old = Normative(
            id=old_id,
            product_code=catalog["product_code"],
            warehouse_code=catalog["warehouse_code"],
            quantity=10,
            unit="шт",
            client_name=name,
            client_id=SYSTEM_CLIENT_ID,
            expiry_date=Request.add_months(date.today(), 6),
            category="A",
        )
        session.add(old)
        await session.commit()

    request_id = await _create_submitted(
        client,
        commercial_token,
        client_name=name,
        expiry_date=valid_expiry_date(),
        items=[
            {
                "product_code": catalog["product_code"],
                "warehouse_code": catalog["warehouse_code"],
                "quantity_requested": 20,
                "unit": "шт",
            }
        ],
    )
    pp = await client.post(
        f"/api/v1/approvals/pp/{request_id}/action",
        headers=auth_header(pp_token),
        json={"action": "approve"},
    )
    assert pp.status_code == 200, pp.text
    economy = await client.post(
        f"/api/v1/approvals/economy/{request_id}/action",
        headers=auth_header(economist_token),
        json={"action": "approve"},
    )
    assert economy.status_code == 200, economy.text

    async with AsyncSessionLocal() as session:
        replaced = await session.get(Normative, old_id)
        assert replaced is not None
        assert replaced.deleted_at is not None
        events = (
            await session.scalars(
                select(Event).where(
                    Event.request_id == request_id,
                    Event.event_type == "normative_replaced_by_name",
                )
            )
        ).all()
        assert len(events) == 1
        assert events[0].payload["old_normative_id"] == str(old_id)
        assert events[0].payload["client_name"] == name

    second_id = await _create_submitted(
        client,
        commercial_token,
        client_id=str(client_id),
        expiry_date=valid_expiry_date(),
        items=[
            {
                "product_code": catalog["product_code"],
                "warehouse_code": catalog["warehouse_code"],
                "quantity_requested": 30,
                "unit": "шт",
            }
        ],
    )
    # client_name was popped when client_id is passed; the helper only adds
    # client_id from client_name. Pass the directory id explicitly above.
    await client.post(
        f"/api/v1/approvals/pp/{second_id}/action",
        headers=auth_header(pp_token),
        json={"action": "approve"},
    )
    second_economy = await client.post(
        f"/api/v1/approvals/economy/{second_id}/action",
        headers=auth_header(economist_token),
        json={"action": "approve"},
    )
    assert second_economy.status_code == 200, second_economy.text
    async with AsyncSessionLocal() as session:
        events = (
            await session.scalars(
                select(Event).where(
                    Event.request_id == second_id,
                    Event.event_type == "normative_replaced_by_name",
                )
            )
        ).all()
        assert events == []
    await delete_request(request_id)
    await delete_request(second_id)
