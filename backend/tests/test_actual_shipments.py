from datetime import date
from decimal import Decimal
from io import BytesIO

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, func, select

from app.core.database import AsyncSessionLocal
from app.models.actual_shipment import ActualShipment
from app.models.client import Client
from app.models.object import Object
from tests.conftest import AuthUser, auth_header, ensure_client, login_token

HEADERS = [
    "Пункт отгрузки",
    "Клиент",
    "Имя клиента",
    "№ продукта",
    "Назв. позиции",
    "Вес нетто",
    "Дата фактического ОМ",
    "Транспортный заказ",
    "Описание Маршрута",
]


def _line(
    client_code,
    product_code,
    weight,
    shipment_date,
    order=None,
    route=None,
    point="2R05",
    client_name="",
    product_name="",
) -> list:
    return [
        point,
        client_code,
        client_name,
        product_code,
        product_name,
        weight,
        shipment_date,
        order,
        route,
    ]


def _xlsx(rows: list[list], headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(content: bytes, name: str = "shipments.xlsx"):
    return {
        "file": (
            name,
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }


async def _client_code(name: str) -> tuple:
    client_id = await ensure_client(name)
    async with AsyncSessionLocal() as session:
        client = await session.get(Client, client_id)
        assert client is not None
        await session.execute(
            delete(ActualShipment).where(ActualShipment.client_id == client_id)
        )
        await session.commit()
        return client_id, int(client.client_code)


async def test_preview_does_not_save_until_confirm(
    client: AsyncClient,
    logistics_user: AuthUser,
    catalog: dict,
) -> None:
    client_id, client_code = await _client_code("Тест фактической реализации")
    content = _xlsx(
        [
            _line(
                client_code,
                catalog["product_code"],
                12.5,
                "01.10.2026",
                "НК-1",
                "Север",
            )
        ]
    )
    token = await login_token(client, logistics_user)
    preview = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        headers=auth_header(token),
        files=_upload(content),
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()["data"]
    assert body["saved"] is False
    assert body["inserted"] == 0
    assert body["valid_rows"] == 1
    assert body["errors"] == 0
    assert body["preview"][0]["client_name"] == "Тест фактической реализации"
    assert body["preview"][0]["route"] == "Север"
    assert body["preview"][0]["shipment_date"] == "2026-10-01"

    async with AsyncSessionLocal() as session:
        stored = await session.scalar(
            select(func.count())
            .select_from(ActualShipment)
            .where(ActualShipment.client_id == client_id)
        )
        assert stored == 0

    confirmed = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        params={"confirm": True},
        headers=auth_header(token),
        files=_upload(content),
    )
    assert confirmed.status_code == 200, confirmed.text
    saved = confirmed.json()["data"]
    assert saved["saved"] is True
    assert saved["inserted"] == 1
    assert saved["message"] == "Сохранено 1, ошибок 0"

    async with AsyncSessionLocal() as session:
        row = await session.scalar(
            select(ActualShipment).where(ActualShipment.client_id == client_id)
        )
        assert row is not None
        assert row.quantity == Decimal("12.50")
        assert row.unit == "кг"
        assert row.loading_point == "2R05"
        assert row.invoice_number == "НК-1"
        assert row.route == "Север"
        assert row.shipment_date == date(2026, 10, 1)
        assert row.uploaded_by == logistics_user.id


async def test_confirm_saves_valid_rows_and_reports_errors(
    client: AsyncClient,
    economist_user: AuthUser,
    catalog: dict,
) -> None:
    client_id, client_code = await _client_code("Тест реализации ошибки")
    content = _xlsx(
        [
            _line(client_code, catalog["product_code"], 2, "15.10.2026"),
            _line(
                999999999,
                catalog["product_code"],
                1,
                "15.10.2026",
                client_name="Чужой клиент",
            ),
            _line(
                client_code,
                99999999,
                1,
                "15.10.2026",
                product_name="Нет такого",
            ),
            _line(client_code, catalog["product_code"], 0, "15.10.2026"),
            _line(client_code, catalog["product_code"], 1, "15.10.2026", "Y" * 26),
            _line(client_code, catalog["product_code"], 3, None),
        ]
    )
    token = await login_token(client, economist_user)
    response = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        params={"confirm": True},
        headers=auth_header(token),
        files=_upload(content),
    )
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["inserted"] == 1
    assert data["valid_rows"] == 1
    assert data["skipped_rows"] == 1
    assert data["errors"] == 2
    assert data["missing_clients"] == [
        {"client_code": 999999999, "client_name": "Чужой клиент", "rows": 1}
    ]
    assert data["missing_products"] == [
        {"product_code": 99999999, "product_name": "Нет такого", "rows": 1}
    ]
    messages = [item["message"] for item in data["error_details"]]
    assert any("25" in item for item in messages)
    assert any("больше 0" in item for item in messages)

    async with AsyncSessionLocal() as session:
        rows = (
            await session.scalars(
                select(ActualShipment).where(ActualShipment.client_id == client_id)
            )
        ).all()
        assert len(rows) == 1
        assert rows[0].unit == "кг"
        assert rows[0].loading_point == "2R05"
        assert rows[0].product_code == catalog["product_code"]


async def test_list_filters_shipments(
    client: AsyncClient,
    pp_user: AuthUser,
    guest_user: AuthUser,
    catalog: dict,
) -> None:
    client_id, client_code = await _client_code("Тест реализации фильтр")
    content = _xlsx(
        [
            _line(
                client_code,
                catalog["product_code"],
                4,
                "01.09.2026",
                "A-1",
                "Юг",
                point="T9X1",
            ),
            _line(
                client_code,
                catalog["product_code_2"],
                8,
                "20.09.2026",
                "A-2",
                "Север",
                point="T9X2",
            ),
        ]
    )
    token = await login_token(client, pp_user)
    uploaded = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        params={"confirm": True},
        headers=auth_header(token),
        files=_upload(content),
    )
    assert uploaded.status_code == 200, uploaded.text
    assert uploaded.json()["data"]["inserted"] == 2

    guest = await login_token(client, guest_user)
    by_product = await client.get(
        "/api/v1/logistics/actual-shipments",
        params={"client_id": str(client_id), "product_code": catalog["product_code"]},
        headers=auth_header(guest),
    )
    assert by_product.status_code == 200, by_product.text
    body = by_product.json()
    assert body["meta"]["total"] == 1
    assert body["data"][0]["invoice_number"] == "A-1"
    assert body["data"][0]["product_name"]

    by_date = await client.get(
        "/api/v1/logistics/actual-shipments",
        params={
            "client_id": str(client_id),
            "date_from": "2026-09-15",
            "date_to": "2026-09-30",
        },
        headers=auth_header(guest),
    )
    assert by_date.status_code == 200, by_date.text
    assert by_date.json()["meta"]["total"] == 1
    assert by_date.json()["data"][0]["route"] == "Север"

    async with AsyncSessionLocal() as session:
        warehouse = await session.get(Object, 9901)
        if warehouse is None:
            session.add(
                Object(
                    code=9901,
                    name="Склад теста отгрузки",
                    city="Тест",
                    type="warehouse",
                    loading_point="T9X1",
                    is_active=True,
                )
            )
        else:
            warehouse.deleted_at = None
            warehouse.loading_point = "T9X1"
            warehouse.is_active = True
            warehouse.type = "warehouse"
        await session.commit()

    by_warehouse = await client.get(
        "/api/v1/logistics/actual-shipments",
        params={"client_id": str(client_id), "warehouse_code": 9901},
        headers=auth_header(guest),
    )
    assert by_warehouse.status_code == 200, by_warehouse.text
    assert by_warehouse.json()["meta"]["total"] == 1
    assert by_warehouse.json()["data"][0]["loading_point"] == "T9X1"

    unknown_warehouse = await client.get(
        "/api/v1/logistics/actual-shipments",
        params={"client_id": str(client_id), "warehouse_code": 990199},
        headers=auth_header(guest),
    )
    assert unknown_warehouse.status_code == 200, unknown_warehouse.text
    assert unknown_warehouse.json()["meta"]["total"] == 0

    async with AsyncSessionLocal() as session:
        await session.execute(delete(Object).where(Object.code == 9901))
        await session.commit()

    invalid = await client.get(
        "/api/v1/logistics/actual-shipments",
        params={"date_from": "2026-10-02", "date_to": "2026-10-01"},
        headers=auth_header(guest),
    )
    assert invalid.status_code == 400


async def test_preview_limits_rows_but_confirm_saves_all(
    client: AsyncClient,
    logistics_user: AuthUser,
    catalog: dict,
) -> None:
    client_id, client_code = await _client_code("Тест реализации 51")
    rows = [
        _line(client_code, catalog["product_code"], 1, "02.10.2026", f"N-{index}", "")
        for index in range(51)
    ]
    content = _xlsx(rows)
    token = await login_token(client, logistics_user)
    preview = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        headers=auth_header(token),
        files=_upload(content),
    )
    assert preview.status_code == 200, preview.text
    data = preview.json()["data"]
    assert data["valid_rows"] == 51
    assert len(data["preview"]) == 50
    assert data["saved"] is False

    confirmed = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        params={"confirm": True},
        headers=auth_header(token),
        files=_upload(content),
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["data"]["inserted"] == 51
    async with AsyncSessionLocal() as session:
        stored = await session.scalar(
            select(func.count())
            .select_from(ActualShipment)
            .where(ActualShipment.client_id == client_id)
        )
        assert stored == 51


async def test_template_and_upload_roles(
    client: AsyncClient,
    pp_user: AuthUser,
    test_user: AuthUser,
) -> None:
    editor = await login_token(client, pp_user)
    template = await client.get(
        "/api/v1/logistics/actual-shipments/template",
        headers=auth_header(editor),
    )
    assert template.status_code == 200, template.text
    workbook = load_workbook(BytesIO(template.content))
    headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
    assert headers[: len(HEADERS)] == HEADERS

    commercial = await login_token(client, test_user)
    forbidden_template = await client.get(
        "/api/v1/logistics/actual-shipments/template",
        headers=auth_header(commercial),
    )
    assert forbidden_template.status_code == 403

    forbidden_upload = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        headers=auth_header(commercial),
        files=_upload(_xlsx([])),
    )
    assert forbidden_upload.status_code == 403

    readable = await client.get(
        "/api/v1/logistics/actual-shipments",
        headers=auth_header(commercial),
    )
    assert readable.status_code == 200


async def test_upload_rejects_non_excel_and_missing_columns(
    client: AsyncClient,
    logistics_user: AuthUser,
) -> None:
    token = await login_token(client, logistics_user)
    rejected = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        headers=auth_header(token),
        files={"file": ("shipments.csv", b"a,b", "text/csv")},
    )
    assert rejected.status_code == 400
    assert rejected.json()["error"]["code"] == "INVALID_FILE"

    missing = await client.post(
        "/api/v1/logistics/actual-shipments/upload",
        headers=auth_header(token),
        files=_upload(_xlsx([], headers=["client_code", "quantity"])),
    )
    assert missing.status_code == 400
    assert "№ продукта" in missing.json()["error"]["message"]
